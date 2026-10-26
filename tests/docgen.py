"""Build small PDF and DOCX fixtures. Uses system fonts: Tahoma and Microsoft YaHei on Windows,
Garuda and WenQuanYi Micro Hei on Ubuntu (fonts-tlwg-garuda-ttf, fonts-wqy-microhei)."""

import ctypes
import sys
from pathlib import Path

import docx
import numpy as np
import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_c

if sys.platform == "win32":
    LATIN_THAI = Path("C:/Windows/Fonts/tahoma.ttf")
    CJK = Path("C:/Windows/Fonts/msyh.ttc")
else:
    LATIN_THAI = Path("/usr/share/fonts/truetype/tlwg/Garuda.ttf")
    CJK = Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc")


def make_pdf(
    path: Path, pages: list[list[str]], font: Path = LATIN_THAI, invisible: bool = False
) -> Path:
    """Write a PDF: one inner list per page, one text line per item (empty list = no text).
    With invisible=True the text is in the text layer but never drawn."""
    pdf = pdfium.PdfDocument.new()
    data = font.read_bytes()
    buffer = (ctypes.c_ubyte * len(data)).from_buffer_copy(data)
    handle = pdfium_c.FPDFText_LoadFont(pdf, buffer, len(data), pdfium_c.FPDF_FONT_TRUETYPE, True)
    if not handle:
        pdf.close()
        raise RuntimeError(f"pdfium could not load {font}")
    try:
        for lines in pages:
            page = pdf.new_page(595, 842)
            y = 800
            for line in lines:
                obj = pdfium_c.FPDFPageObj_CreateTextObj(pdf, handle, 12.0)
                # FPDF_WIDESTRING is UTF-16; create_unicode_buffer is UTF-32 on Linux.
                utf16 = (line + "\0").encode("utf-16-le")
                text = (ctypes.c_ushort * (len(utf16) // 2)).from_buffer_copy(utf16)
                pdfium_c.FPDFText_SetText(
                    obj, ctypes.cast(text, ctypes.POINTER(pdfium_c.FPDF_WCHAR))
                )
                if invisible:  # text in the text layer that is never drawn
                    pdfium_c.FPDFTextObj_SetTextRenderMode(
                        obj, pdfium_c.FPDF_TEXTRENDERMODE_INVISIBLE
                    )
                pdfium_c.FPDFPageObj_Transform(obj, 1, 0, 0, 1, 50, y)
                pdfium_c.FPDFPage_InsertObject(page, obj)
                y -= 20
            pdfium_c.FPDFPage_GenerateContent(page)
            page.close()
        path.parent.mkdir(parents=True, exist_ok=True)
        pdf.save(path)
    finally:
        pdfium_c.FPDFFont_Close(handle)
        pdf.close()
    return path


def image_pdf(src: Path, dest: Path, dpi: int = 200) -> Path:
    """A scan of src: every page drawn as a grayscale image at `dpi`, with no text layer."""
    doc = pdfium.PdfDocument(src)
    out = pdfium.PdfDocument.new()
    try:
        for index in range(len(doc)):
            page = doc[index]
            width, height = page.get_size()
            gray = np.ascontiguousarray(page.render(scale=dpi / 72, grayscale=True).to_numpy())
            if gray.ndim == 3:
                gray = gray[:, :, 0]
            bitmap = pdfium.PdfBitmap.new_native(
                gray.shape[1], gray.shape[0], pdfium_c.FPDFBitmap_Gray
            )
            bitmap.to_numpy()[...] = gray
            target = out.new_page(width, height)
            image = pdfium.PdfImage.new(out)
            image.set_bitmap(bitmap)
            image.set_matrix(pdfium.PdfMatrix().scale(width, height))
            target.insert_obj(image)
            target.gen_content()
        dest.parent.mkdir(parents=True, exist_ok=True)
        out.save(dest)
    finally:
        out.close()
        doc.close()
    return dest


def make_docx(path: Path, blocks: list[tuple[int, str] | tuple[int, list[list[str]]]]) -> Path:
    """Write a DOCX from (heading level, text) pairs; level 0 is a body paragraph.

    A block whose second item is a list of rows (lists of cell texts) becomes a table.
    """
    document = docx.Document()
    for level, text in blocks:
        if isinstance(text, list):
            table = document.add_table(rows=len(text), cols=len(text[0]))
            for row, texts in zip(table.rows, text, strict=True):
                for cell, value in zip(row.cells, texts, strict=True):
                    cell.text = value
        elif level:
            document.add_heading(text, level=level)
        else:
            document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path
