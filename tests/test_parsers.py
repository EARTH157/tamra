import codecs
from pathlib import Path

import pytest
from docgen import CJK, LATIN_THAI, image_pdf, make_docx, make_pdf

from tamra.ingest.parsers import (
    ECHO_MIN,
    ParseError,
    Unit,
    decode_text,
    location,
    parse_file,
    parse_pdf,
    thai_echoes,
    with_ocr,
)

FIXTURES = Path(__file__).parent / "fixtures" / "ocr"


@pytest.fixture(autouse=True)
def _fonts():
    if not (LATIN_THAI.exists() and CJK.exists()):
        pytest.skip("Windows fonts Tahoma and Microsoft YaHei are needed to build PDFs")


def test_pdf_pages_become_units_with_page_numbers(tmp_path):
    path = make_pdf(
        tmp_path / "t.pdf",
        [["Lease agreement", "สัญญาเช่าบ้านมีอายุสามปี"], ["Page two English line."]],
    )
    doc = parse_file(path)
    assert (doc.kind, doc.note) == ("pdf", None)
    assert [u.page for u in doc.units] == [1, 2]
    assert "สัญญาเช่าบ้านมีอายุสามปี" in doc.units[0].text
    assert doc.units[1].text.strip() == "Page two English line."


def test_chinese_pdf_text_is_extracted(tmp_path):
    path = make_pdf(tmp_path / "z.pdf", [["租赁期限为三年。"]], font=CJK)
    assert parse_file(path).units[0].text.strip() == "租赁期限为三年。"


def test_pages_without_text_are_reported(tmp_path):
    doc = parse_file(make_pdf(tmp_path / "s.pdf", [["text page"], []]))
    assert [u.page for u in doc.units] == [1]
    assert doc.note == "no text layer on page(s) 2 (needs OCR)"


def test_a_broken_pdf_raises_parse_error(tmp_path):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"%PDF-1.7 this is not a real pdf")
    with pytest.raises(ParseError, match="cannot open PDF"):
        parse_file(bad)


def test_docx_paragraphs_carry_their_heading_path(tmp_path):
    path = make_docx(
        tmp_path / "d.docx",
        [
            (1, "บทที่ 1"),
            (0, "ย่อหน้าแรก"),
            (2, "Section 1.1"),
            (0, "Body"),
            (0, "   "),
            (1, "Chapter 2"),
            (0, "End"),
        ],
    )
    doc = parse_file(path)
    assert doc.kind == "docx"
    assert [(u.text, u.heading_path, u.paragraph) for u in doc.units] == [
        ("บทที่ 1", ("บทที่ 1",), 0),
        ("ย่อหน้าแรก", ("บทที่ 1",), 1),
        ("Section 1.1", ("บทที่ 1", "Section 1.1"), 2),
        ("Body", ("บทที่ 1", "Section 1.1"), 3),
        ("Chapter 2", ("Chapter 2",), 5),
        ("End", ("Chapter 2",), 6),
    ]


def test_docx_units_know_their_own_heading_level(tmp_path):
    path = make_docx(
        tmp_path / "h.docx",
        [(1, "Chapter"), (0, "Chapter"), (2, "Section"), (0, "Body")],
    )
    doc = parse_file(path)
    assert [u.heading_level for u in doc.units] == [1, 0, 2, 0]
    assert Unit("plain", line=1).heading_level == 0


def test_a_broken_docx_raises_parse_error(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"PK not really a zip")
    with pytest.raises(ParseError, match="cannot open DOCX"):
        parse_file(bad)


def test_text_paragraphs_have_line_numbers(tmp_path):
    path = tmp_path / "n.md"
    path.write_bytes(b"# Title\r\n\r\nFirst line\r\nsecond line\r\n\r\n\r\nLast")
    doc = parse_file(path)
    assert doc.kind == "text"
    assert [(u.text, u.line) for u in doc.units] == [
        ("# Title", 1),
        ("First line\nsecond line", 3),
        ("Last", 7),
    ]


@pytest.mark.parametrize(
    "data",
    [
        "สัญญาเช่าบ้าน".encode(),
        codecs.BOM_UTF8 + "สัญญาเช่าบ้าน".encode(),
        "สัญญาเช่าบ้าน".encode("utf-16"),
        "สัญญาเช่าบ้าน".encode("tis-620"),
    ],
)
def test_decode_text_handles_utf_and_short_thai_legacy_files(data):
    assert decode_text(data) == "สัญญาเช่าบ้าน"


def test_decode_text_does_not_mistake_chinese_gbk_for_thai():
    text = (
        "卡诺家具产品保修条款：沙发框架保修五年，布料和海绵保修两年，餐桌和椅子保修一年。"
        "保修期内的维修服务免费，技术人员会在三个工作日内上门检查。"
    )
    assert decode_text(text.encode("gbk")) == text


def test_unsupported_types_raise_parse_error(tmp_path):
    path = tmp_path / "x.csv"
    path.write_text("a,b", encoding="utf-8")
    with pytest.raises(ParseError, match="unsupported"):
        parse_file(path)


def test_location_shapes():
    pdf_a, pdf_b = Unit("abcdef", page=2, char=10), Unit("ghij", page=3)
    assert location("pdf", pdf_a, 1, pdf_b, 4) == {
        "kind": "pdf",
        "page_start": 2,
        "char_start": 11,
        "page_end": 3,
        "char_end": 4,
    }
    doc_a = Unit("x", paragraph=4, heading_path=("H", "S"))
    doc_b = Unit("y", paragraph=9, heading_path=("H", "T"))
    assert location("docx", doc_a, 0, doc_b, 1) == {
        "kind": "docx",
        "heading_path": ["H", "S"],
        "paragraph_start": 4,
        "paragraph_end": 9,
    }
    text = Unit("one\ntwo\nthree", line=10)
    assert location("text", text, 4, text, 13) == {"kind": "text", "line_start": 11, "line_end": 12}


def test_docx_table_rows_become_units_in_document_order(tmp_path):
    path = make_docx(
        tmp_path / "t.docx",
        [
            (1, "Fees"),
            (0, "Before the table."),
            (0, [["Item", "Price", "Note"], ["Rent", "", "monthly"]]),
            (0, "After the table."),
        ],
    )
    doc = parse_file(path)
    assert [(u.text, u.paragraph, u.heading_path, u.heading_level) for u in doc.units] == [
        ("Fees", 0, ("Fees",), 1),
        ("Before the table.", 1, ("Fees",), 0),
        ("Item | Price | Note", 2, ("Fees",), 0),
        ("Rent | monthly", 3, ("Fees",), 0),
        ("After the table.", 4, ("Fees",), 0),
    ]


def test_docx_merged_cells_are_read_once_and_nested_tables_follow_their_row(tmp_path):
    import docx

    document = docx.Document()
    title = document.add_table(rows=1, cols=3)
    title.cell(0, 0).merge(title.cell(0, 2)).text = "Title"  # horizontal merge
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Group"  # vertical merge
    table.cell(0, 1).text = "one"
    table.cell(1, 1).text = "two"
    table.cell(2, 0).text = "Outer"
    inner = table.cell(2, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "nested"
    inner.cell(0, 1).text = "row"
    path = tmp_path / "m.docx"
    document.save(str(path))
    assert [(u.text, u.paragraph) for u in parse_file(path).units] == [
        ("Title", 0),
        ("Group | one", 1),
        ("two", 2),
        ("Outer", 3),
        ("nested | row", 4),
    ]


CLEAN_THAI = [
    "ระเบียบการใช้ห้องปฏิบัติการ",
    "ผู้ใช้ห้องปฏิบัติการทุกคนต้องผ่านการอบรมความปลอดภัยก่อนเริ่มใช้งานเครื่องมือ",
    "ห้ามนำอาหารและเครื่องดื่มเข้ามาในห้องปฏิบัติการ",
]


def test_thai_echoes_count_the_fragments_a_skia_print_repeats():
    assert thai_echoes("บอร์ดร์ พัฒพั") == 2  # inline: a cluster again after its word
    assert thai_echoes("สรุป\nรุ\nผลการทดลอง") == 1  # a short line that repeats the tail above
    assert thai_echoes("\u0e34abc") == 1  # a mark with no base letter
    assert thai_echoes("\n".join(CLEAN_THAI)) == 0 < ECHO_MIN == 3
    for truth in sorted(FIXTURES.glob("th_*.txt")):
        assert thai_echoes(truth.read_text(encoding="utf-8")) == 0, truth.name


def test_a_page_without_text_is_left_for_ocr(tmp_path):
    doc = parse_pdf(make_pdf(tmp_path / "a.pdf", [["Page one text here."], []]))
    assert [u.page for u in doc.units] == [1]
    assert doc.ocr_pages == (2,)
    assert doc.note == "no text layer on page(s) 2 (needs OCR)"


@pytest.mark.parametrize("name", [f"th_{i}.pdf" for i in range(6)])
def test_a_browser_printed_thai_layer_is_sent_to_ocr_but_kept(name):
    doc = parse_pdf(FIXTURES / name)
    assert doc.ocr_pages == (1,)
    assert [u.page for u in doc.units] == [1] and doc.note is None


def test_clean_text_layers_are_not_sent_to_ocr(tmp_path):
    thai = parse_pdf(make_pdf(tmp_path / "th.pdf", [CLEAN_THAI * 3]))
    english = parse_pdf(make_pdf(tmp_path / "en.pdf", [["The lease term is three years."] * 20]))
    chinese = parse_pdf(
        make_pdf(tmp_path / "zh.pdf", [["租期为三年，押金为两个月租金。"] * 10], CJK)
    )
    assert thai.ocr_pages == english.ocr_pages == chinese.ocr_pages == ()


def test_text_that_is_not_drawn_is_sent_to_ocr(tmp_path):
    lines = ["Hidden text that the reader never sees on this page."] * 4  # over 100 characters
    hidden = parse_pdf(make_pdf(tmp_path / "h.pdf", [lines], invisible=True))
    short = parse_pdf(make_pdf(tmp_path / "s.pdf", [["Hidden footer."]], invisible=True))
    assert hidden.ocr_pages == (1,) and short.ocr_pages == ()


def test_with_ocr_replaces_pages_and_drops_empty_ones(tmp_path):
    doc = parse_pdf(make_pdf(tmp_path / "a.pdf", [["One."], ["Two."], []]))
    read = with_ocr(doc, {2: "Two, read again.", 3: "Three.", 1: ""})
    assert [(u.page, u.text, u.ocr) for u in read.units] == [
        (2, "Two, read again.", True),
        (3, "Three.", True),
    ]
    assert read.kind == "pdf" and read.note == doc.note


def test_image_pdf_has_no_text_layer(tmp_path):
    scan = image_pdf(make_pdf(tmp_path / "a.pdf", [["Visible words."]]), tmp_path / "scan.pdf")
    doc = parse_pdf(scan)
    assert doc.units == [] and doc.ocr_pages == (1,)
