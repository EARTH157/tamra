"""Parse PDF, DOCX, TXT, and MD files into text units that remember where they came from."""

import codecs
import re
import threading
from collections.abc import Iterator
from dataclasses import dataclass, replace
from pathlib import Path

import docx
import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_c
from charset_normalizer import from_bytes
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

SUPPORTED = frozenset({".pdf", ".docx", ".txt", ".md"})

# PDFium is not thread-safe: indexing runs in a worker thread while the viewer serves requests
# on others, so every open-to-close use of a PdfDocument holds this lock.
PDFIUM_LOCK = threading.RLock()

_P, _TBL = qn("w:p"), qn("w:tbl")
_HEADING = re.compile(r"Heading (\d)")
_THAI_CONSONANTS = range(0x0E01, 0x0E2F)
_THAI_MARKS = frozenset([0x0E31, *range(0x0E34, 0x0E3B), *range(0x0E47, 0x0E4F)])

# Pages that OCR reads although they have a text layer (spec §5). Measured on the owner's files
# and the spike's Edge prints: see docs/spikes/2026-10-m5-results.md.
ECHO_MIN = 3  # garbled Thai: Chromium/Skia prints have 5-44 echoes a page, clean Thai 0
INK_SCALE = 0.25  # an 18 dpi render tells a drawn page from an empty one
INVISIBLE_MIN_CHARS = 100
INVISIBLE_INK_PER_CHAR = 0.1  # dark pixels per character: drawn pages have 0.35-8, hidden 0
SCAN_IMAGE_AREA = 0.5  # one image this share of the page is a scan; logos and photos are less

_THAI_LETTER = "\u0e01-\u0e2e"
_THAI_MARK = "\u0e31\u0e34-\u0e3a\u0e47-\u0e4e"
_ECHO = re.compile(rf"([{_THAI_LETTER}][{_THAI_MARK}]+)[{_THAI_LETTER}]\1")
_ORPHAN_MARK = re.compile(rf"(?<![\u0e01-\u0e5b])[{_THAI_MARK}]")


class ParseError(Exception):
    """The file could not be read as a document of its type."""


@dataclass(frozen=True)
class Unit:
    """A run of document text and the coordinates of its first character.

    pdf: page (1-based) and char (offset of text[0] in the page's extracted text).
    docx: paragraph (0-based index of the body paragraph or table row), heading_path,
    and heading_level.
    text: line (1-based number of the line holding text[0]).
    """

    text: str
    page: int = 0
    char: int = 0
    paragraph: int = 0
    heading_path: tuple[str, ...] = ()
    line: int = 0
    heading_level: int = 0  # docx: 1-9 when the paragraph is itself a heading or title
    ocr: bool = False  # pdf: the text was read by OCR, not taken from the text layer


@dataclass(frozen=True)
class ParsedDoc:
    kind: str  # "pdf" | "docx" | "text"
    units: list[Unit]
    note: str | None = None  # shown next to the file, e.g. pages without a text layer
    ocr_pages: tuple[int, ...] = ()  # pdf: pages OCR should read (1-based)


def location(kind: str, first: Unit, first_offset: int, last: Unit, last_offset: int) -> dict:
    """Where the text from first.text[first_offset] to last.text[last_offset - 1] sits."""
    if kind == "pdf":
        where = {
            "kind": "pdf",
            "page_start": first.page,
            "char_start": first.char + first_offset,
            "page_end": last.page,
            "char_end": last.char + last_offset,
        }
        if first.ocr or last.ocr:
            where["ocr"] = True  # the viewer takes this page's boxes from the OCR store
        return where
    if kind == "docx":
        return {
            "kind": "docx",
            "heading_path": list(first.heading_path),
            "paragraph_start": first.paragraph,
            "paragraph_end": last.paragraph,
        }
    return {
        "kind": "text",
        "line_start": first.line + first.text.count("\n", 0, first_offset),
        "line_end": last.line + last.text.count("\n", 0, max(last_offset - 1, 0)),
    }


def parse_file(path: Path) -> ParsedDoc:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix in (".txt", ".md"):
        return parse_text(path)
    raise ParseError(f"unsupported file type: {path.suffix}")


def parse_pdf(path: Path) -> ParsedDoc:
    """One unit per page with a text layer; pages without one are listed in the note. Pages
    without a text layer, or with one that is garbled or not drawn, are listed in ocr_pages."""
    with PDFIUM_LOCK:
        return _parse_pdf(path)


def _parse_pdf(path: Path) -> ParsedDoc:
    try:
        pdf = pdfium.PdfDocument(path)
    except pdfium.PdfiumError as e:
        raise ParseError(f"cannot open PDF: {e}") from e
    units: list[Unit] = []
    empty: list[int] = []
    broken: list[int] = []
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_range()
                finally:
                    textpage.close()
                unreadable = bool(text.strip()) and (
                    thai_echoes(text) >= ECHO_MIN or _not_drawn(page, text)
                )
            finally:
                page.close()
            if text.strip():
                units.append(Unit(text=text, page=index + 1))
            else:
                empty.append(index + 1)
            if unreadable:
                broken.append(index + 1)
    finally:
        pdf.close()
    note = None
    if empty:
        note = "no text layer on page(s) " + ", ".join(map(str, empty)) + " (needs OCR)"
    return ParsedDoc("pdf", units, note, ocr_pages=tuple(sorted(empty + broken)))


def thai_echoes(text: str) -> int:
    """Signs of a Chromium/Skia-printed Thai text layer: a cluster repeated after its word
    ("บอร์ดร์"), a short line repeating the tail of the line above ("สรุป\\nรุ"), and marks with
    no base letter."""
    count = len(_ECHO.findall(text)) + len(_ORPHAN_MARK.findall(text))
    previous = ""
    for line in text.splitlines():
        line = line.strip()
        short = line and len(line) <= 3 and "\u0e01" <= line[0] <= "\u0e2e"
        # Skia repeats a tail fragment; a whole line again is a table's repeated cell
        if short and line in previous[-4:] and line != previous:
            count += 1
        if line:
            previous = line
    return count


def _not_drawn(page, text: str) -> bool:
    """A page with plenty of text whose render is almost blank: its text is invisible."""
    chars = len(text.strip())
    if chars < INVISIBLE_MIN_CHARS:
        return False
    if _is_scan(page):  # the invisible text is the image's OCR layer: skip the slow render
        return False
    bitmap = page.render(scale=INK_SCALE, grayscale=True)
    try:
        dark = int((bitmap.to_numpy() < 224).sum())
    finally:
        bitmap.close()
    return dark / chars < INVISIBLE_INK_PER_CHAR


def _is_scan(page) -> bool:
    """One image object covers at least SCAN_IMAGE_AREA of the page."""
    width, height = page.get_size()
    for obj in page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE], max_depth=2):
        left, bottom, right, top = obj.get_bounds()  # FPDFPageObj_GetBounds, in points
        if (right - left) * (top - bottom) >= SCAN_IMAGE_AREA * width * height:
            return True
    return False


def with_ocr(doc: ParsedDoc, pages: dict[int, str]) -> ParsedDoc:
    """doc with the text of these pages replaced by OCR text; where OCR found no text, the
    page keeps its text-layer unit, if it has one."""
    units = [u for u in doc.units if not pages.get(u.page, "").strip()]
    units += [Unit(text=t, page=p, ocr=True) for p, t in pages.items() if t.strip()]
    return replace(doc, units=sorted(units, key=lambda u: u.page))


def _table_rows(table: DocxTable) -> Iterator[str]:
    """The text of each row ("" when empty), then the rows of tables nested in its cells."""
    seen: set = set()  # a merged cell is reported once, in its first row
    for row in table.rows:
        cells = []
        for cell in row.cells:  # a horizontal merge returns the same cell once per grid column
            if cell._tc not in seen:
                seen.add(cell._tc)
                cells.append(cell)
        yield " | ".join(t for t in (cell.text.strip() for cell in cells) if t)
        for cell in cells:
            for nested in cell.tables:
                yield from _table_rows(nested)


def parse_docx(path: Path) -> ParsedDoc:
    """One unit per non-empty paragraph and per table row, in document order.

    A row is its cell texts joined by " | " (empty cells skipped, a merged cell once); rows of
    nested tables follow their row. `paragraph` counts every body paragraph and table row, so it
    is a stable position in the document.
    """
    try:
        document = docx.Document(str(path))
    except Exception as e:  # python-docx raises several types for a broken package
        raise ParseError(f"cannot open DOCX: {e}") from e
    units: list[Unit] = []
    headings: list[str] = []
    index = 0
    for child in document.element.body.iterchildren():
        if child.tag == _TBL:
            for text in _table_rows(DocxTable(child, document)):
                if text:
                    units.append(Unit(text=text, paragraph=index, heading_path=tuple(headings)))
                index += 1
        elif child.tag == _P:
            paragraph = DocxParagraph(child, document)
            text = paragraph.text
            if text.strip():
                style = paragraph.style.name if paragraph.style is not None else ""
                match = _HEADING.fullmatch(style)
                level = int(match.group(1)) if match else (1 if style == "Title" else 0)
                if level:
                    headings = headings[: level - 1] + [text.strip()]
                units.append(
                    Unit(
                        text=text,
                        paragraph=index,
                        heading_path=tuple(headings),
                        heading_level=level,
                    )
                )
            index += 1
    return ParsedDoc("docx", units)


def parse_text(path: Path) -> ParsedDoc:
    """One unit per paragraph (a run of non-blank lines), with its first line number."""
    try:
        data = path.read_bytes()
    except OSError as e:
        raise ParseError(f"cannot read file: {e}") from e
    text = decode_text(data).replace("\r\n", "\n").replace("\r", "\n")
    units: list[Unit] = []
    block: list[str] = []
    start = 0
    for number, line in enumerate(text.split("\n"), start=1):
        if line.strip():
            if not block:
                start = number
            block.append(line)
        elif block:
            units.append(Unit(text="\n".join(block), line=start))
            block = []
    if block:
        units.append(Unit(text="\n".join(block), line=start))
    return ParsedDoc("text", units)


def decode_text(data: bytes) -> str:
    """Decode a text file: BOMs, then UTF-8, then Thai cp874 (TIS-620), then a detected codec."""
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8) :].decode("utf-8", errors="replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        thai = data.decode("cp874")
    except UnicodeDecodeError:
        thai = None
    if thai is not None and _reads_as_thai(thai):
        return thai
    best = from_bytes(data).best()
    return str(best) if best is not None else data.decode("latin-1")


def _reads_as_thai(text: str) -> bool:
    """Mostly Thai letters, with vowel and tone marks sitting on consonants.

    Other legacy encodings decoded as cp874 (GBK Chinese, for example) also produce Thai
    letters, but their marks land after arbitrary characters, so the second test fails.
    """
    letters = [c for c in text if c.isalpha()]
    thai = [c for c in letters if 0x0E00 <= ord(c) <= 0x0E7F]
    if not letters or len(thai) < 0.3 * len(letters):
        return False
    marks = placed = 0
    previous = 0
    for c in text:
        code = ord(c)
        if code in _THAI_MARKS:
            marks += 1
            if previous in _THAI_CONSONANTS or previous in _THAI_MARKS:
                placed += 1
        previous = code
    return marks == 0 or placed >= 0.95 * marks
